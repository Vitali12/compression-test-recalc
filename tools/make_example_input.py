#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Generate examples/example_input.xlsx — a synthetic input file in the
format expected by recalc_compression.py.

Synthetic model used so the results are easy to verify:

- the testing machine (specimen 0, "without specimen") has a stiffness of
  200 N/mm, i.e. its recorded travel is u12 = P / 200;
- specimen 1 has a specimen stiffness K3 = 1000 N/mm,
  specimen 2 has K3 = 1400 N/mm;
- the recorded total travel of a specimen is u = P/200 + P/K3.

The program must therefore return K3 = 1000 / 1400 N/mm, K = 2*K3,
sigma = Pmax / A and E = K * L / A.
"""
from pathlib import Path

from openpyxl import Workbook

K_MACHINE = 200.0                 # N/mm — machine stiffness (specimen 0)
K3 = {0: None, 1: 1000.0, 2: 1400.0}   # N/mm — specimen stiffness
L = {0: None, 1: 33.6, 2: 34.0}        # mm — specimen height
A = {0: None, 1: 113.1, 2: 113.4}      # mm² — contact area
N_POINTS = 601                    # P = 0 … 1200 N, step 2 N
P_MAX = 1200.0                    # N


def main(out: str = "examples/example_input.xlsx") -> None:
    wb = Workbook()
    ws = wb.active

    ws["A1"] = "K-27"                                   # row 1 — series label

    for j in range(3):                                  # blocks in columns A / D / G
        c0 = 1 + 3 * j
        spec = j                                        # specimen number: 0, 1, 2

        ws.cell(row=2, column=c0, value=spec)           # row 2 — specimen No. (0 = no specimen)
        if L[spec]:
            ws.cell(row=3, column=c0, value=L[spec])    # row 3 — height L, mm
        if A[spec]:
            ws.cell(row=4, column=c0, value=A[spec])    # row 4 — contact area A, mm²

        ws.cell(row=5, column=c0, value="Время")        # row 5 — headers
        ws.cell(row=5, column=c0 + 1, value="Сила")
        ws.cell(row=5, column=c0 + 2, value="Ход")

        ws.cell(row=6, column=c0, value="сек")          # row 6 — units (optional)
        ws.cell(row=6, column=c0 + 1, value="Н")
        ws.cell(row=6, column=c0 + 2, value="мм")

        for i in range(N_POINTS):                       # rows 7+ — data
            p = i * P_MAX / (N_POINTS - 1)
            u = p / K_MACHINE + (p / K3[spec] if K3[spec] else 0.0)
            r = 7 + i
            ws.cell(row=r, column=c0, value=round(0.1 * i, 2))   # time, s
            ws.cell(row=r, column=c0 + 1, value=round(p, 2))     # force, N
            ws.cell(row=r, column=c0 + 2, value=round(u, 6))     # travel, mm

    path = Path(out)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print("saved", path)


if __name__ == "__main__":
    main()
