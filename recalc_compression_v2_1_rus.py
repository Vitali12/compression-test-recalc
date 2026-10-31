#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Пересчёт данных испытания на сжатие с нейтрализацией податливости машины.

Формат исходного листа Excel:
    строка 1  — № серии (например, К-27)
    строка 2  — номер образца: 1, 2, 3, ...; 0 — испытание без образца
    строка 3  — контактная площадь A, мм²
    строка 4  — заголовки: Время | Сила | Ход (на каждый образец по 3 столбца)
    строка 5  — единицы: сек | Н | мм
    далее     — данные

Количество образцов не ограничено (5, 10 и больше). Образец 0 — эталонная
кривая «без образца», по ней учитывается податливость испытательной машины.

Запуск без аргументов открывает графический интерфейс.

Примеры CLI:
    python recalc_compression.py файл.xlsx --length 33.6
    python recalc_compression.py файл.xlsx --samples 1,2,3 --out результат
"""

from __future__ import annotations

import argparse
import math
import os
import re
import sys
import traceback
from dataclasses import dataclass

import numpy as np
import pandas as pd

import matplotlib

matplotlib.use("Agg")
from matplotlib.figure import Figure

DEFAULT_LENGTH = 33.6
DEFAULT_M = 250
DEFAULT_MACHINE_SMOOTH = 101

NAN = float("nan")

# --------------------------------------------------------------------------
# Вспомогательные функции
# --------------------------------------------------------------------------


def norm_text(value) -> str:
    """Нормализация текста заголовков: нижний регистр, без лишних пробелов."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    s = str(value).replace("\xa0", " ").replace("\u2009", " ").replace("\u202f", " ")
    s = re.sub(r"\s+", " ", s).strip().lower().replace("ё", "е")
    return s


_NUM_FULL = re.compile(r"^[-+]?\d+(?:[.,]\d+)?(?:[eE][-+]?\d+)?$")


def parse_number(value) -> float:
    """Число из ячейки Excel, в том числе «1 165,8» (неразрывные пробелы, запятая)."""
    if value is None:
        return NAN
    if isinstance(value, bool):
        return NAN
    if isinstance(value, (int, float, np.integer, np.floating)):
        v = float(value)
        return v if math.isfinite(v) else NAN
    s = str(value)
    s = s.replace("\xa0", "").replace("\u2009", "").replace("\u202f", "")
    s = s.replace(" ", "").replace(",", ".").strip()
    if not _NUM_FULL.match(s):
        return NAN
    try:
        return float(s)
    except ValueError:
        return NAN


def numeric_array(series: pd.Series) -> np.ndarray:
    """Столбец в float с поддержкой запятой как десятичного разделителя."""
    if pd.api.types.is_numeric_dtype(series):
        return series.to_numpy(dtype=float, na_value=np.nan)
    s = series.astype(str)
    for bad in ("\xa0", "\u2009", "\u202f", " "):
        s = s.str.replace(bad, "", regex=False)
    s = s.str.replace(",", ".", regex=False)
    return pd.to_numeric(s, errors="coerce").to_numpy(dtype=float)


def fmt(value, digits=1, dash="—") -> str:
    if value is None:
        return dash
    try:
        v = float(value)
    except (TypeError, ValueError):
        return dash
    if not math.isfinite(v):
        return dash
    return f"{v:.{digits}f}"


def fmt_int(value, dash="—") -> str:
    if value is None:
        return dash
    try:
        v = float(value)
    except (TypeError, ValueError):
        return dash
    if not math.isfinite(v):
        return dash
    return f"{v:,.0f}".replace(",", " ")


# --------------------------------------------------------------------------
# Чтение файла
# --------------------------------------------------------------------------


@dataclass
class SampleBlock:
    label: str
    area: float
    col: int
    series: str = ""
    length: float = NAN


@dataclass
class Dataset:
    path: str
    sheet: str
    series: str
    header_row: int
    blocks: list
    data: dict
    machine_label: str | None
    warnings: list


def _find_header_row(df: pd.DataFrame, max_scan: int = 80) -> int | None:
    n = min(len(df), max_scan)
    for r in range(n):
        row = df.iloc[r]
        for v in row:
            if norm_text(v) == "время":
                return r
    return None


def _area_from_row(df: pd.DataFrame, row: int, col: int) -> float:
    if row is None or row < 0:
        return NAN
    v = parse_number(df.iat[row, col])
    return v


def load_dataset(path: str, sheet=0) -> Dataset:
    """Читает лист Excel и находит все блоки «Время | Сила | Ход»."""
    xl = pd.ExcelFile(path)
    if isinstance(sheet, int):
        if sheet >= len(xl.sheet_names):
            raise ValueError(f"В файле нет листа с индексом {sheet}")
        sheet_name = xl.sheet_names[sheet]
    else:
        if sheet not in xl.sheet_names:
            raise ValueError(f"Лист «{sheet}» не найден. Доступны: {', '.join(xl.sheet_names)}")
        sheet_name = sheet

    df = pd.read_excel(path, sheet_name=sheet_name, header=None)
    if df.shape[1] < 3:
        raise ValueError("В листе меньше трёх столбцов — не найдены данные.")

    warnings = []
    header_row = _find_header_row(df)
    if header_row is None:
        raise ValueError(
            "Не найден заголовок «Время». Ожидается формат: строка 1 — № серии, "
            "строка 2 — № образца, строка 3 — площадь контакта, строка 4 — "
            "Время | Сила | Ход."
        )

    sample_row = None
    area_row = None
    series_row = None

    # --- разбор шапки: номер образца, высота L, площадь A ---
    probe_cols = [c for c in range(df.shape[1] - 2)
                  if norm_text(df.iat[header_row, c]) == "время"]

    def row_fill(r):
        vals = [parse_number(df.iat[r, c]) for c in probe_cols]
        return sum(1 for v in vals if math.isfinite(v))

    numeric_rows = []
    for r in range(max(1, header_row - 6), header_row):
        if row_fill(r) >= max(1, int(0.6 * len(probe_cols))):
            numeric_rows.append(r)
    numeric_rows.sort()

    length_row = None
    if len(numeric_rows) >= 3:
        sample_row, length_row, area_row = numeric_rows[-3], numeric_rows[-2], numeric_rows[-1]
    elif len(numeric_rows) == 2:
        sample_row, area_row = numeric_rows[-2], numeric_rows[-1]
    elif len(numeric_rows) == 1:
        area_row = numeric_rows[0]

    if sample_row is not None:
        ints = 0
        total = 0
        for c in probe_cols:
            v = parse_number(df.iat[sample_row, c])
            if math.isfinite(v):
                total += 1
                if abs(v - round(v)) < 1e-9 and 0 <= v <= 999:
                    ints += 1
        if total and ints < max(1, int(0.6 * total)):
            warnings.append(f"Строка {sample_row + 1} не похожа на номера образцов — номера присвоены по порядку.")
            sample_row = None

    if sample_row is not None and sample_row > 0:
        series_row = sample_row - 1
    else:
        series_row = 0

    blocks: list[SampleBlock] = []
    seen: dict[str, int] = {}
    used_cols = set()

    for c in range(df.shape[1] - 2):
        if norm_text(df.iat[header_row, c]) != "время":
            continue
        if c in used_cols:
            continue
        h_force = norm_text(df.iat[header_row, c + 1])
        h_disp = norm_text(df.iat[header_row, c + 2])
        if not (h_force.startswith("сил") and h_disp.startswith("ход")):
            warnings.append(
                f"Столбец {c + 1}: после «Время» ожидались «Сила» и «Ход», "
                f"найдено «{h_force}» и «{h_disp}» — блок пропущен."
            )
            continue

        raw_label = df.iat[sample_row, c] if sample_row is not None else None
        label = ""
        num = parse_number(raw_label)
        if math.isfinite(num) and abs(num - round(num)) < 1e-9:
            label = str(int(round(num)))
        elif raw_label is not None and norm_text(raw_label):
            label = str(raw_label).strip()
        if not label:
            label = str(len(blocks) + 1)
            warnings.append(f"У блока в столбце {c + 1} не указан номер — присвоен №{label}.")

        if label in seen:
            seen[label] += 1
            new_label = f"{label} ({seen[label]})"
            warnings.append(f"Номер образца {label} повторяется — блок переименован в «{new_label}».")
            label = new_label
        else:
            seen[label] = 1

        area = _area_from_row(df, area_row, c) if area_row is not None else NAN
        if not math.isfinite(area):
            # ищем строку с площадью выше заголовка (до 4 строк)
            for rr in range((area_row if area_row is not None else header_row) - 1,
                            max(-1, header_row - 5), -1):
                v = _area_from_row(df, rr, c)
                if math.isfinite(v) and v > 0:
                    area = v
                    break

        length = _area_from_row(df, length_row, c) if length_row is not None else NAN
        if not (math.isfinite(length) and length > 0):
            length = NAN

        series = ""
        if series_row is not None:
            sv = df.iat[series_row, c]
            if sv is not None and norm_text(sv):
                series = str(sv).strip()

        blocks.append(SampleBlock(label=label, area=area if math.isfinite(area) else NAN,
                                  col=c, series=series, length=length))
        used_cols.update((c, c + 1, c + 2))

    if not blocks:
        raise ValueError("Не найдено ни одного блока «Время | Сила | Ход».")

    # определяем первую строку данных (строка единиц может отсутствовать)
    units_row = header_row + 1
    first_data = units_row + 1
    if units_row < len(df):
        probe = df.iat[units_row, blocks[0].col]
        if math.isfinite(parse_number(probe)):
            first_data = units_row

    data = {}
    for b in blocks:
        t = numeric_array(df.iloc[first_data:, b.col])
        p = numeric_array(df.iloc[first_data:, b.col + 1])
        u = numeric_array(df.iloc[first_data:, b.col + 2])
        mask = np.isfinite(t) & np.isfinite(p) & np.isfinite(u)
        t, p, u = t[mask], p[mask], u[mask]
        if t.size < 10:
            warnings.append(f"Образец {b.label}: всего {t.size} строк с данными — слишком мало.")
        data[b.label] = (t, p, u)

    machine_label = None
    for b in blocks:
        if b.label.strip() == "0":
            machine_label = b.label
            break
    if machine_label is None:
        warnings.append("Образец 0 («без образца») не найден — податливость машины учесть нельзя.")

    if length_row is None:
        warnings.append("Высота образца L в шапке не найдена — используется значение по умолчанию.")
    else:
        missing_l = [b.label for b in blocks if b.label != machine_label and not math.isfinite(b.length)]
        if missing_l:
            warnings.append("Не указана высота L для образцов: " + ", ".join(missing_l) + ".")

    series_name = ""
    for b in blocks:
        if b.series:
            series_name = b.series
            break
    if not series_name:
        series_name = os.path.splitext(os.path.basename(path))[0]

    return Dataset(
        path=path,
        sheet=sheet_name,
        series=series_name,
        header_row=header_row,
        blocks=blocks,
        data=data,
        machine_label=machine_label,
        warnings=warnings,
    )


# --------------------------------------------------------------------------
# Расчёт
# --------------------------------------------------------------------------


@dataclass
class SampleResult:
    label: str
    area: float = NAN
    length: float = NAN
    ok: bool = False
    error: str = ""
    note: str = ""
    pmax: float = NAN
    sigma: float = NAN
    p123: float = NAN
    u123: float = NAN
    u12: float = NAN
    u3: float = NAN
    k3: float = NAN
    k_corr: float = NAN
    k123: float = NAN
    k12: float = NAN
    k123_corr: float = NAN
    k12_corr: float = NAN
    e_mpa: float = NAN
    i_star: int = -1
    peak: int = -1
    p_curve: np.ndarray | None = None
    u123_curve: np.ndarray | None = None
    u12_curve: np.ndarray | None = None
    u3_curve: np.ndarray | None = None
    ui: np.ndarray | None = None
    ki: np.ndarray | None = None


def smooth(y: np.ndarray, w: int) -> np.ndarray:
    w = int(w)
    if w < 2:
        return y.copy()
    if w % 2 == 0:
        w += 1
    pad = w // 2
    ypad = np.pad(y, pad, mode="edge")
    return np.convolve(ypad, np.ones(w) / w, mode="valid")


def prep_machine(p: np.ndarray, u: np.ndarray, peak: int, smooth_w: int):
    p = p[:peak + 1]
    u = u[:peak + 1]
    order = np.argsort(p)
    p, u = p[order], u[order]
    u = smooth(u, smooth_w)
    u = np.maximum.accumulate(u)
    keep = np.concatenate(([True], np.diff(p) > 0))
    return p[keep], u[keep]


def stiffness(p: np.ndarray, u: np.ndarray, m: int):
    m = max(1, int(m))
    n = p.size
    if n <= 2 * m:
        raise ValueError(f"окно m={m} слишком велико для {n} точек")
    i = np.arange(m, n - m)
    du = u[i + m] - u[i - m]
    dp = p[i + m] - p[i - m]
    k = np.full(n, np.nan)
    with np.errstate(divide="ignore", invalid="ignore"):
        k[m:n - m] = np.where(du != 0, dp / du, np.nan)
    return i, k


def compute_sample(p, u, machine, length, area, m, machine_smooth, label="") -> SampleResult:
    p = np.asarray(p, dtype=float)
    u = np.asarray(u, dtype=float)
    r = SampleResult(label=label, area=area, length=length)

    if p.size < 10:
        r.error = f"мало точек ({p.size})"
        return r

    r.peak = int(np.argmax(p))
    r.pmax = float(p[r.peak])

    try:
        i, k = stiffness(p, u, m)
    except ValueError as exc:
        r.error = str(exc)
        return r

    valid = i < r.peak
    i_valid = i[valid]
    k_valid = k[i_valid]
    finite = np.isfinite(k_valid)
    if not finite.any():
        r.error = "не удалось вычислить жёсткость (нулевые приращения хода)"
        return r
    i_star = int(i_valid[finite][int(np.argmax(k_valid[finite]))])

    r.i_star = i_star
    r.p123 = float(p[i_star])
    r.u123 = float(u[i_star])

    if machine is not None:
        xp12, fp12 = machine["xp"], machine["fp"]
        r.u12 = float(np.interp(r.p123, xp12, fp12))
        r.u12_curve = np.interp(p[:r.peak + 1], xp12, fp12)
        if r.p123 > float(xp12[-1]):
            r.note = "P123 выше диапазона эталонной кривой — экстраполяция"
    else:
        r.u12 = 0.0
        r.u12_curve = np.zeros(r.peak + 1, dtype=float)

    r.u3 = r.u123 - r.u12
    if not math.isfinite(r.u3) or r.u3 <= 0:
        r.error = "деформация образца u3 ≤ 0 — проверьте кривую «без образца»"
        return r

    r.k3 = r.p123 / r.u3
    r.k_corr = 2.0 * r.k3
    r.k123 = r.p123 / r.u123 if r.u123 else NAN
    r.k12 = r.p123 / r.u12 if r.u12 else NAN
    r.k123_corr = 2.0 * r.k123
    r.k12_corr = 2.0 * r.k12

    if math.isfinite(area) and area > 0:
        r.sigma = r.pmax / area
        r.e_mpa = r.k_corr * length / area
    else:
        r.note = (r.note + "; " if r.note else "") + "не указана площадь контакта"

    r.p_curve = p[:r.peak + 1]
    r.u123_curve = u[:r.peak + 1]
    r.u3_curve = r.u123_curve - r.u12_curve
    r.ui = u[i]
    r.ki = k[i]
    r.ok = True
    return r


def compute_all(ds: Dataset, lengths: dict, m: int, machine_smooth: int,
                area_fallback: float = NAN, use_machine: bool = True):
    """Считает все образцы. Возвращает (results, machine_info)."""
    machine_info = None
    machine = None

    if use_machine and ds.machine_label is not None:
        t0, p0, u0 = ds.data[ds.machine_label]
        if p0.size >= 10:
            peak0 = int(np.argmax(p0))
            xp12, fp12 = prep_machine(p0, u0, peak0, machine_smooth)
            machine = {"xp": xp12, "fp": fp12}
            machine_info = {
                "label": ds.machine_label,
                "pmax": float(p0.max()),
                "xp": xp12,
                "fp": fp12,
                "p": p0,
                "u": u0,
            }

    results: list[SampleResult] = []
    for b in ds.blocks:
        if b.label == ds.machine_label:
            continue
        t, p, u = ds.data[b.label]
        area = b.area if math.isfinite(b.area) else area_fallback
        length = lengths.get(b.label)
        if length is None or not math.isfinite(float(length)) or float(length) <= 0:
            length = b.length if (math.isfinite(b.length) and b.length > 0) else DEFAULT_LENGTH
        r = compute_sample(p, u, machine, float(length), area, m, machine_smooth, label=b.label)
        results.append(r)

    return results, machine_info


def stats(values) -> tuple[float, float, float, int]:
    a = np.array([v for v in values if v is not None and math.isfinite(float(v))], dtype=float)
    if a.size == 0:
        return NAN, NAN, NAN, 0
    mean = float(a.mean())
    std = float(a.std(ddof=1)) if a.size > 1 else 0.0
    cv = float(std / mean * 100.0) if mean else NAN
    return mean, std, cv, int(a.size)


# --------------------------------------------------------------------------
# Текстовая сводка
# --------------------------------------------------------------------------


def summary_dataframe(results: list[SampleResult]) -> pd.DataFrame:
    rows = []
    for r in results:
        rows.append({
            "№": r.label,
            "A, мм²": fmt(r.area, 1),
            "Pmax, Н": f"{r.pmax:.0f}",
            "σ, МПа": fmt(r.sigma, 1),
            "P123, Н": f"{r.p123:.0f}",
            "u123, мм": fmt(r.u123, 4),
            "u12, мм": fmt(r.u12, 4),
            "u3, мм": fmt(r.u3, 4),
            "K3, Н/мм": f"{r.k3:.0f}",
            "K=2·K3, Н/мм": f"{r.k_corr:.0f}",
            "E, ГПа": fmt(r.e_mpa / 1000.0 if math.isfinite(r.e_mpa) else NAN, 2),
        })
    return pd.DataFrame(rows)


def build_summary_text(ds: Dataset, results: list[SampleResult],
                       selected: list[SampleResult], params: dict) -> str:
    lines = []
    lines.append(f"Серия: {ds.series}")
    lines.append(f"Файл:  {ds.path}")
    lines.append(f"Лист:  {ds.sheet}")
    n_total = len(results)
    extra = " + образец 0 (без образца)" if ds.machine_label is not None else ""
    lines.append(f"Образцов в файле: {n_total}{extra}")
    lines.append(f"Выбрано: {len(selected)} из {n_total}")
    lines.append("")
    lines.append("Параметры расчёта")
    lens = [r.length for r in selected if math.isfinite(r.length)]
    if lens:
        if max(lens) - min(lens) < 1e-9:
            lines.append(f"  Длина образца L = {lens[0]:.2f} мм")
        else:
            lines.append(f"  Длина образца L = {min(lens):.2f}–{max(lens):.2f} мм (по образцам)")
    else:
        lines.append(f"  Длина образца L = {params['length']:.2f} мм")
    lines.append(f"  Окно жёсткости m, точек = {params['m']}")
    lines.append(f"  Сглаживание кривой машины, точек = {params['machine_smooth']}")
    lines.append(f"  Податливость машины: {'учитывается' if params['use_machine'] else 'не учитывается'}")
    lines.append("")

    if not selected:
        lines.append("Образцы не выбраны.")
        return "\n".join(lines)

    lines.append("Результаты по выбранным образцам")
    lines.append(summary_dataframe(selected).to_string(index=False))
    lines.append("")

    s_mean, s_std, s_cv, s_n = stats([r.sigma for r in selected])
    e_mean, e_std, e_cv, e_n = stats([r.e_mpa for r in selected])
    k_mean, k_std, k_cv, k_n = stats([r.k_corr for r in selected])
    k3_mean, k3_std, k3_cv, _ = stats([r.k3 for r in selected])

    lines.append(f"Средние значения по выбранным образцам (n = {s_n})")
    if s_n:
        lines.append(f"  Предел прочности σ   = {s_mean:.2f} ± {s_std:.2f} МПа     (CV {s_cv:.1f} %)")
        lines.append(f"  Модуль Юнга E        = {e_mean / 1000.0:.2f} ± {e_std / 1000.0:.2f} ГПа"
                     f"     (CV {e_cv:.1f} %)")
        lines.append(f"  Показатель K = 2·K3  = {fmt_int(k_mean)} ± {fmt_int(k_std)} Н/мм"
                     f"   (CV {k_cv:.1f} %)")
        lines.append(f"  K3                   = {fmt_int(k3_mean)} ± {fmt_int(k3_std)} Н/мм"
                     f"   (CV {k3_cv:.1f} %)")
    else:
        lines.append("  нет достоверных данных")

    notes = []
    for r in selected:
        if r.note:
            notes.append(f"  №{r.label}: {r.note}")
    if notes:
        lines.append("")
        lines.append("Примечания")
        lines.extend(notes)

    return "\n".join(lines)


# --------------------------------------------------------------------------
# Графики
# --------------------------------------------------------------------------


def build_file_info(ds: Dataset) -> str:
    lines = [f"Серия: {ds.series}", f"Файл:  {ds.path}", f"Лист:  {ds.sheet}", "", "Обнаруженные образцы:"]
    for b in ds.blocks:
        tag = "  — без образца (эталон)" if b.label == ds.machine_label else ""
        area = f"A = {b.area:.1f} мм²" if math.isfinite(b.area) else "A не указана"
        length = f"L = {b.length:.2f} мм" if math.isfinite(b.length) else "L не указана"
        n = len(ds.data[b.label][0])
        lines.append(f"  №{b.label}: {n} точек, {area}, {length}{tag}")
    if ds.warnings:
        lines.append("")
        lines.append("Предупреждения:")
        for w in ds.warnings:
            lines.append(f"  - {w}")
    return "\n".join(lines)


def figure_overview(results: list[SampleResult], machine_info: dict | None) -> Figure:
    fig = Figure(figsize=(11.4, 7.4))
    ax1 = fig.add_subplot(2, 2, 1)
    ax2 = fig.add_subplot(2, 2, 2)
    ax3 = fig.add_subplot(2, 2, 3)
    ax4 = fig.add_subplot(2, 2, 4)

    for r in results:
        if not r.ok:
            continue
        ax1.plot(r.u123_curve, r.p_curve / 1000.0, lw=1.1, label=f"№{r.label}")
        ax1.plot(r.u123, r.p123 / 1000.0, "o", ms=4, color="black", zorder=5)
    if machine_info is not None:
        ax1.plot(machine_info["fp"], machine_info["xp"] / 1000.0, "k--", lw=1.0,
                 label="без образца (0)")
    ax1.set_xlabel("Ход u123, мм")
    ax1.set_ylabel("Сила P, кН")
    ax1.set_title("Кривые нагружения")
    ax1.grid(True, alpha=0.3)
    ax1.legend(fontsize=8, ncol=2)

    labels = [f"№{r.label}" for r in results]
    x = np.arange(len(results))

    def bar(ax, values, title, ylabel, color, scale=1.0, fmt_digits=1):
        vals = np.array([v * scale if math.isfinite(v) else np.nan for v in values], dtype=float)
        ax.bar(x, vals, color=color, alpha=0.85, edgecolor="white")
        if np.isfinite(vals).any():
            mean = float(np.nanmean(vals))
            ax.axhline(mean, color="#c0392b", ls="--", lw=1.2,
                       label=f"среднее {mean:.{fmt_digits}f}")
            ax.legend(fontsize=8)
            for xi, v in zip(x, vals):
                if math.isfinite(v):
                    ax.text(xi, v, f"{v:.{fmt_digits}f}", ha="center", va="bottom", fontsize=7)
        ax.set_xticks(x)
        ax.set_xticklabels(labels, fontsize=8)
        ax.set_title(title)
        ax.set_ylabel(ylabel)
        ax.grid(True, axis="y", alpha=0.3)

    bar(ax2, [r.sigma for r in results], "Предел прочности", "σ, МПа", "#4c78a8", 1.0, 1)
    bar(ax3, [r.e_mpa / 1000.0 for r in results], "Модуль Юнга", "E, ГПа", "#59a14f", 1.0, 1)
    bar(ax4, [r.k_corr / 1000.0 for r in results], "Показатель K = 2·K3", "K, кН/мм", "#f28e2b", 1.0, 0)

    fig.tight_layout()
    return fig


def figure_sample(r: SampleResult) -> Figure:
    fig = Figure(figsize=(12.4, 4.2))
    ax1, ax2, ax3 = fig.subplots(1, 3)

    ax1.plot(r.u123_curve, r.p_curve / 1000.0, "g", lw=0.8, label="с образцом (123)")
    if r.u12_curve is not None:
        ax1.plot(r.u12_curve, r.p_curve / 1000.0, "r", lw=0.8, label="без образца (12)")
    ax1.plot(r.u3_curve, r.p_curve / 1000.0, "b", lw=1.2, label="деформация образца (3)")
    ax1.plot(r.u123, r.p123 / 1000.0, "ko", ms=6, label=f"P123={r.p123 / 1000:.1f} кН")
    ax1.plot(r.u12, r.p123 / 1000.0, "ks", ms=6)
    ax1.set_xlabel("Перемещение u, мм")
    ax1.set_ylabel("Сила P, кН")
    ax1.legend(fontsize=7)
    ax1.set_title(f"Образец {r.label}: кривые")
    ax1.grid(True, alpha=0.3)

    if r.ui is not None:
        ax2.plot(r.ui, r.ki, "g", lw=0.5)
        ax2.plot(r.u123, r.k123, "ko", ms=7)
    ax2.set_xlabel("u123, мм")
    ax2.set_ylabel("K123, Н/мм")
    ax2.set_title("Жёсткость системы 123")
    ax2.grid(True, alpha=0.3)

    ax3.plot(r.u3_curve, r.p_curve / 1000.0, "b", lw=1.0)
    ax3.plot(r.u3, r.p123 / 1000.0, "ko", ms=7)
    ax3.set_xlabel("u3, мм")
    ax3.set_ylabel("P, кН")
    parts = []
    if math.isfinite(r.e_mpa):
        parts.append(f"E={r.e_mpa / 1000:.1f} ГПа")
    if math.isfinite(r.sigma):
        parts.append(f"σ={r.sigma:.1f} МПа")
    ax3.set_title("Образец: " + ", ".join(parts) if parts else f"Образец {r.label}")
    ax3.grid(True, alpha=0.3)

    fig.tight_layout()
    return fig


# --------------------------------------------------------------------------
# Экспорт в Excel
# --------------------------------------------------------------------------


def _safe_sheet_name(name: str, existing: set) -> str:
    clean = re.sub(r"[\\/*?:\[\]]", "_", str(name)).strip() or "Лист"
    clean = clean[:31]
    base = clean
    i = 2
    while clean in existing:
        suffix = f"_{i}"
        clean = base[:31 - len(suffix)] + suffix
        i += 1
    existing.add(clean)
    return clean


def export_excel(path: str, ds: Dataset, results: list[SampleResult],
                 selected: list[SampleResult], params: dict,
                 machine_info: dict | None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"

    bold = Font(bold=True)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5597")
    sub_fill = PatternFill("solid", fgColor="DDEBF7")
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A1"] = "Результаты испытания на сжатие"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:E1")

    info = [
        ("Серия", ds.series),
        ("Файл", ds.path),
        ("Лист", ds.sheet),
        ("Образцов", len(results)),
        ("Податливость машины", "учитывается" if params["use_machine"] else "не учитывается"),
        ("L (медиана), мм" if any(math.isfinite(b.length) and b.length > 0
                                   for b in ds.blocks if b.label != ds.machine_label) else "L, мм",
         params["length"]),
        ("m, точек", params["m"]),
        ("Сглаживание машины, точек", params["machine_smooth"]),
    ]
    row = 3
    for key, value in info:
        ws.cell(row=row, column=1, value=key).font = bold
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    header_row = row

    headers = [
        "№ образца", "A, мм²", "L, мм", "Pmax, Н", "σ, МПа", "P123, Н",
        "u123, мм", "u12, мм", "u3, мм", "K123, Н/мм", "K12, Н/мм",
        "K3, Н/мм", "K=2·K3, Н/мм", "K123 испр., Н/мм", "K12 испр., Н/мм",
        "E, МПа", "E, ГПа", "i*", "Пик (индекс)", "Примечание",
    ]
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=j, value=h)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[header_row].height = 34

    row = header_row + 1
    for r in selected:
        values = [
            r.label, r.area, r.length, r.pmax, r.sigma, r.p123,
            r.u123, r.u12, r.u3, r.k123, r.k12,
            r.k3, r.k_corr, r.k123_corr, r.k12_corr,
            r.e_mpa, r.e_mpa / 1000.0 if math.isfinite(r.e_mpa) else NAN,
            r.i_star if r.i_star >= 0 else None,
            r.peak if r.peak >= 0 else None,
            r.note,
        ]
        for j, v in enumerate(values, start=1):
            if isinstance(v, float) and not math.isfinite(v):
                v = None
            cell = ws.cell(row=row, column=j, value=v)
            cell.border = border
            if j in (1, 18, 19):
                cell.alignment = Alignment(horizontal="center")
        row += 1

    for j in range(1, len(headers) + 1):
        ws.cell(row=row, column=j).border = Border(top=Side(style="medium"))

    row += 2
    ws.cell(row=row, column=1, value="Средние значения по выбранным образцам").font = Font(bold=True, size=12)
    row += 1
    mean_headers = ["Показатель", "Среднее", "Ст. откл.", "CV, %", "n"]
    for j, h in enumerate(mean_headers, start=1):
        cell = ws.cell(row=row, column=j, value=h)
        cell.font = bold
        cell.fill = sub_fill
        cell.border = border
        cell.alignment = center
    row += 1

    s_mean, s_std, s_cv, s_n = stats([r.sigma for r in selected])
    e_mean, e_std, e_cv, _ = stats([r.e_mpa for r in selected])
    k_mean, k_std, k_cv, _ = stats([r.k_corr for r in selected])
    k3_mean, k3_std, k3_cv, _ = stats([r.k3 for r in selected])

    mean_rows = [
        ("Предел прочности σ, МПа", s_mean, s_std, s_cv, s_n),
        ("Модуль Юнга E, ГПа", e_mean / 1000.0 if math.isfinite(e_mean) else NAN,
         e_std / 1000.0 if math.isfinite(e_std) else NAN, e_cv, s_n),
        ("Показатель K = 2·K3, Н/мм", k_mean, k_std, k_cv, s_n),
        ("K3, Н/мм", k3_mean, k3_std, k3_cv, s_n),
    ]
    for name, mean, std, cv, n in mean_rows:
        ws.cell(row=row, column=1, value=name).border = border
        for j, v in enumerate((mean, std, cv, n), start=2):
            cell = ws.cell(row=row, column=j, value=None if isinstance(v, float) and not math.isfinite(v) else v)
            cell.border = border
        row += 1

    widths = [12, 11, 9, 12, 10, 12, 10, 10, 10, 13, 13, 13, 14, 15, 15, 12, 10, 8, 13, 40]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)

    formats = {
        2: "0.00", 3: "0.00", 4: "#,##0", 5: "0.00", 6: "#,##0",
        7: "0.0000", 8: "0.0000", 9: "0.0000",
        10: "#,##0", 11: "#,##0", 12: "#,##0", 13: "#,##0",
        14: "#,##0", 15: "#,##0", 16: "#,##0", 17: "0.00",
    }
    for rr in range(header_row + 1, header_row + 1 + len(selected)):
        for col, f in formats.items():
            ws.cell(row=rr, column=col).number_format = f

    # лист с параметрами и предупреждениями
    ws2 = wb.create_sheet("Параметры")
    ws2["A1"] = "Параметры расчёта"
    ws2["A1"].font = Font(bold=True, size=12)
    ws2["A2"] = "Длина образца по умолчанию, мм"
    ws2["B2"] = params["length"]
    ws2["A3"] = "Окно жёсткости m, точек"
    ws2["B3"] = params["m"]
    ws2["A4"] = "Сглаживание кривой машины, точек"
    ws2["B4"] = params["machine_smooth"]
    ws2["A5"] = "Податливость машины"
    ws2["B5"] = "учитывается" if params["use_machine"] else "не учитывается"
    ws2["A6"] = "Образец 0"
    ws2["B6"] = "найден" if ds.machine_label is not None else "не найден"
    ws2["A7"] = "Источник L"
    has_file_l = any(math.isfinite(b.length) and b.length > 0 for b in ds.blocks if b.label != ds.machine_label)
    ws2["B7"] = "шапка файла (по образцам)" if has_file_l else "задана вручную"
    ws2.column_dimensions["A"].width = 38
    ws2.column_dimensions["B"].width = 50

    r0 = 8
    if ds.warnings:
        ws2.cell(row=r0, column=1, value="Предупреждения").font = bold
        r0 += 1
        for w in ds.warnings:
            ws2.cell(row=r0, column=1, value=w)
            r0 += 1

    # отдельный лист на каждый выбранный образец
    existing = {"Сводка", "Параметры"}
    for r in selected:
        name = _safe_sheet_name(f"Образец {r.label}", existing)
        sh = wb.create_sheet(name)
        for j, h in enumerate(("P, Н", "u3, мм", "u123, мм", "u12, мм"), start=1):
            cell = sh.cell(row=1, column=j, value=h)
            cell.font = bold
            cell.fill = sub_fill
            cell.border = border
        p = r.p_curve
        u3 = r.u3_curve
        u123 = r.u123_curve
        u12 = r.u12_curve
        for i in range(len(p)):
            sh.cell(row=i + 2, column=1, value=float(p[i]))
            sh.cell(row=i + 2, column=2, value=float(u3[i]) if u3 is not None else None)
            sh.cell(row=i + 2, column=3, value=float(u123[i]) if u123 is not None else None)
            sh.cell(row=i + 2, column=4, value=float(u12[i]) if u12 is not None else None)
        for j, w in enumerate((14, 12, 12, 12), start=1):
            sh.column_dimensions[get_column_letter(j)].width = w
        sh.freeze_panes = "A2"

    wb.save(path)


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------


def set_console_utf8():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def main_cli():
    set_console_utf8()
    ap = argparse.ArgumentParser(
        description="Расчёт испытания на сжатие по серии образцов (образец 0 — без образца)"
    )
    ap.add_argument("xlsx", help="исходный файл Excel")
    ap.add_argument("--sheet", default=0, help="лист: имя или индекс (по умолчанию первый)")
    ap.add_argument("--length", "-L", type=float, default=None,
                    help=f"длина (высота) образца, мм; если не указана — берётся из шапки файла (иначе {DEFAULT_LENGTH})")
    ap.add_argument("--m", type=int, default=DEFAULT_M,
                    help=f"окно расчёта жёсткости, точек (по умолчанию {DEFAULT_M})")
    ap.add_argument("--machine-smooth", type=int, default=DEFAULT_MACHINE_SMOOTH,
                    help=f"сглаживание кривой машины, точек (по умолчанию {DEFAULT_MACHINE_SMOOTH})")
    ap.add_argument("--samples", default="all",
                    help="номера образцов через запятую или all (по умолчанию all)")
    ap.add_argument("--area", type=float, default=None,
                    help="площадь контакта, мм² (если в файле не указана; иначе берётся из файла)")
    ap.add_argument("--no-machine", action="store_true",
                    help="не учитывать податливость машины (игнорировать образец 0)")
    ap.add_argument("--out", default=None,
                    help="префикс выходных файлов (по умолчанию — имя исходного файла)")
    args = ap.parse_args()

    sheet = args.sheet
    if isinstance(sheet, str) and sheet.isdigit():
        sheet = int(sheet)

    ds = load_dataset(args.xlsx, sheet)
    if args.length is not None:
        lengths = {b.label: args.length for b in ds.blocks}
        length_default = args.length
    else:
        lengths = {b.label: (b.length if math.isfinite(b.length) and b.length > 0 else DEFAULT_LENGTH)
                   for b in ds.blocks}
        vals = [b.length for b in ds.blocks
                if b.label != ds.machine_label and math.isfinite(b.length) and b.length > 0]
        length_default = float(np.median(vals)) if vals else DEFAULT_LENGTH
    area_fb = args.area if args.area is not None else NAN

    results, machine_info = compute_all(
        ds, lengths, args.m, args.machine_smooth,
        area_fallback=area_fb, use_machine=not args.no_machine,
    )

    by_label = {r.label: r for r in results}
    if args.samples.strip().lower() == "all":
        selected = [r for r in results if r.ok]
    else:
        wanted = [s.strip() for s in args.samples.split(",") if s.strip()]
        selected = [by_label[s] for s in wanted if s in by_label and by_label[s].ok]
        missing = [s for s in wanted if s not in by_label]
        if missing:
            print(f"Нет таких образцов: {', '.join(missing)}")

    params = {
        "length": length_default,
        "m": args.m,
        "machine_smooth": args.machine_smooth,
        "use_machine": not args.no_machine,
    }

    print(build_summary_text(ds, results, selected, params))

    if not selected:
        print("\nНет рассчитанных образцов для выгрузки.")

    if ds.warnings:
        print("\nПредупреждения:")
        for w in ds.warnings:
            print("  -", w)

    out = args.out
    if not out:
        base = re.sub(r'[\\/:*?"<>|]+', "_", ds.series).strip()
        if not base:
            base = os.path.splitext(os.path.basename(args.xlsx))[0]
        out = os.path.join(os.path.dirname(os.path.abspath(args.xlsx)), base)
    xlsx_path = out + "_результаты.xlsx"
    png_path = out + "_графики.png"

    export_excel(xlsx_path, ds, results, selected, params, machine_info)
    fig = figure_overview(selected, machine_info)
    fig.savefig(png_path, dpi=150)

    print(f"\nРезультаты: {xlsx_path}")
    print(f"Графики:    {png_path}")


# --------------------------------------------------------------------------
# GUI
# --------------------------------------------------------------------------


def main_gui():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

    class App:
        def __init__(self, root):
            self.root = root
            self.ds = None
            self.results = []
            self.by_label = {}
            self.machine_info = None
            self.selected = set()
            self.lengths = {}
            self.iid_to_label = {}
            self._editing = None

            root.title("Расчёт испытания на сжатие — серия образцов")
            try:
                k = float(root.tk.call("tk", "scaling")) or 1.0
            except Exception:
                k = 1.0
            sw, sh = root.winfo_screenwidth(), root.winfo_screenheight()
            w = int(min(1220, max(900, (sw - 60) / k)))
            h = int(min(800, max(600, (sh - 100) / k)))
            root.geometry(f"{w}x{h}+20+10")
            root.minsize(960, 640)
            self._build()

        # ---------------- интерфейс ----------------

        def _build(self):
            top = ttk.LabelFrame(self.root, text="Исходные данные", padding=8)
            top.pack(fill="x", padx=8, pady=(8, 4))

            ttk.Label(top, text="Файл:").grid(row=0, column=0, sticky="w")
            self.path_var = tk.StringVar()
            ttk.Entry(top, textvariable=self.path_var, width=60).grid(
                row=0, column=1, columnspan=3, sticky="we", padx=4)
            ttk.Button(top, text="Обзор…", command=self.browse).grid(row=0, column=4, padx=4)

            ttk.Label(top, text="Лист:").grid(row=1, column=0, sticky="w", pady=(6, 0))
            self.sheet_var = tk.StringVar()
            self.sheet_box = ttk.Combobox(top, textvariable=self.sheet_var, width=24, state="readonly")
            self.sheet_box.grid(row=1, column=1, sticky="w", padx=4, pady=(6, 0))
            self.sheet_box.bind("<<ComboboxSelected>>", lambda e: self.load())
            ttk.Button(top, text="Загрузить", command=self.load).grid(row=1, column=4, padx=4, pady=(6, 0))
            self.file_info = ttk.Label(top, text="Файл не загружен", foreground="#666666")
            self.file_info.grid(row=1, column=2, columnspan=2, sticky="w", pady=(6, 0))
            top.columnconfigure(1, weight=1)
            top.columnconfigure(2, weight=1)

            par = ttk.LabelFrame(self.root, text="Параметры расчёта", padding=8)
            par.pack(fill="x", padx=8, pady=4)

            ttk.Label(par, text="L, мм:").grid(row=0, column=0, sticky="w")
            self.e_len = ttk.Entry(par, width=8)
            self.e_len.grid(row=0, column=1, sticky="w", padx=(2, 12))
            self.e_len.insert(0, f"{DEFAULT_LENGTH:g}")

            ttk.Label(par, text="m, точек:").grid(row=0, column=2, sticky="w")
            self.e_m = ttk.Entry(par, width=6)
            self.e_m.grid(row=0, column=3, sticky="w", padx=(2, 12))
            self.e_m.insert(0, str(DEFAULT_M))

            ttk.Label(par, text="Сглаживание машины:").grid(row=0, column=4, sticky="w")
            self.e_ms = ttk.Entry(par, width=6)
            self.e_ms.grid(row=0, column=5, sticky="w", padx=(2, 12))
            self.e_ms.insert(0, str(DEFAULT_MACHINE_SMOOTH))

            self.var_machine = tk.BooleanVar(value=True)
            ttk.Checkbutton(par, text="Учитывать податливость машины (образец 0)",
                            variable=self.var_machine).grid(row=0, column=6, sticky="w", padx=(0, 16))

            ttk.Label(par, text="A, мм² (если нет в файле):").grid(row=1, column=0, columnspan=2,
                                                              sticky="w", pady=(6, 0))
            self.e_area = ttk.Entry(par, width=9)
            self.e_area.grid(row=1, column=2, sticky="w", padx=(2, 0), pady=(6, 0))

            bottom = ttk.Frame(self.root)
            bottom.pack(side="bottom", fill="x", padx=8, pady=(4, 8))
            self.btn_calc = ttk.Button(bottom, text="Рассчитать", command=self.recalc)
            self.btn_calc.pack(side="left")
            self.btn_xlsx = ttk.Button(bottom, text="Экспорт Excel…", command=self.export_xlsx,
                                       state="disabled")
            self.btn_xlsx.pack(side="left", padx=6)
            self.btn_png = ttk.Button(bottom, text="Экспорт графиков…", command=self.export_png,
                                      state="disabled")
            self.btn_png.pack(side="left")
            self.status = ttk.Label(bottom, text="Загрузите файл Excel", foreground="#666666")
            self.status.pack(side="right")

            body = ttk.Frame(self.root)
            body.pack(fill="both", expand=True, padx=8, pady=4)

            left = ttk.LabelFrame(body, text="Образцы", padding=6)
            left.pack(side="left", fill="y")
            right = ttk.Frame(body)
            right.pack(side="left", fill="both", expand=True, padx=(8, 0))

            cols = ("sel", "num", "area", "length", "pmax", "sigma", "egpa", "k", "note")
            self.tree = ttk.Treeview(left, columns=cols, show="headings", height=12,
                                    selectmode="browse")
            heads = {
                "sel": "✓", "num": "№", "area": "A, мм²", "length": "L, мм",
                "pmax": "Pmax, Н", "sigma": "σ, МПа", "egpa": "E, ГПа",
                "k": "K, Н/мм", "note": "Примечание",
            }
            widths = {
                "sel": 26, "num": 32, "area": 54, "length": 46, "pmax": 60,
                "sigma": 52, "egpa": 52, "k": 74, "note": 110,
            }
            for c in cols:
                self.tree.heading(c, text=heads[c])
                self.tree.column(c, width=widths[c], anchor="center",
                                 stretch=(c == "note"))
            self.tree.tag_configure("machine", foreground="#808080")
            self.tree.tag_configure("error", foreground="#b00020")
            self.tree.grid(row=0, column=0, sticky="nsew")
            sb = ttk.Scrollbar(left, orient="vertical", command=self.tree.yview)
            sb.grid(row=0, column=1, sticky="ns")
            self.tree.configure(yscrollcommand=sb.set)

            self.tree.bind("<Button-1>", self.on_click)
            self.tree.bind("<Double-1>", self.on_double)
            self.tree.bind("<<TreeviewSelect>>", lambda e: self.on_select_row())
            self.tree.bind("<space>", lambda e: self.toggle_many())

            bfrm = ttk.Frame(left)
            bfrm.grid(row=1, column=0, columnspan=2, sticky="we", pady=(6, 0))
            ttk.Button(bfrm, text="Выбрать все", command=self.select_all).pack(side="left")
            ttk.Button(bfrm, text="Снять все", command=self.select_none).pack(side="left", padx=4)
            ttk.Label(left, text="Клик по ✓ — показать/скрыть; двойной клик по L — изменить",
                      foreground="#666666").grid(row=2, column=0, columnspan=2, sticky="w", pady=(4, 0))
            left.rowconfigure(0, weight=1)
            left.columnconfigure(0, weight=1)

            self.nb = ttk.Notebook(right)
            self.nb.pack(fill="both", expand=True)

            tab_sum = ttk.Frame(self.nb)
            self.nb.add(tab_sum, text="Сводка")
            self.txt = tk.Text(tab_sum, wrap="none", width=40, height=12,
                               font=("Consolas", 9), state="disabled")
            ysb = ttk.Scrollbar(tab_sum, orient="vertical", command=self.txt.yview)
            xsb = ttk.Scrollbar(tab_sum, orient="horizontal", command=self.txt.xview)
            self.txt.configure(yscrollcommand=ysb.set, xscrollcommand=xsb.set)
            self.txt.grid(row=0, column=0, sticky="nsew")
            ysb.grid(row=0, column=1, sticky="ns")
            xsb.grid(row=1, column=0, sticky="we")
            tab_sum.rowconfigure(0, weight=1)
            tab_sum.columnconfigure(0, weight=1)

            self.tab_over = ttk.Frame(self.nb)
            self.nb.add(self.tab_over, text="Общие графики")
            self.tab_detail = ttk.Frame(self.nb)
            self.nb.add(self.tab_detail, text="Подробно по образцу")

        # ---------------- загрузка ----------------

        def browse(self):
            path = filedialog.askopenfilename(
                title="Выберите файл данных",
                filetypes=[("Excel", "*.xlsx *.xlsm"), ("Все файлы", "*.*")])
            if not path:
                return
            self.path_var.set(path)
            try:
                names = pd.ExcelFile(path).sheet_names
            except Exception as exc:
                messagebox.showerror("Ошибка", f"Не удалось открыть файл:\n{exc}")
                return
            self.sheet_box["values"] = names
            if names:
                self.sheet_box.current(0)
            self.load()

        def load(self):
            path = self.path_var.get().strip()
            if not path:
                return
            sheet = self.sheet_var.get() or 0
            self.root.config(cursor="watch")
            self.status.configure(text="Чтение файла…")
            self.root.update_idletasks()
            try:
                ds = load_dataset(path, sheet)
            except Exception as exc:
                self.root.config(cursor="")
                messagebox.showerror("Ошибка чтения", str(exc))
                self.status.configure(text="Ошибка чтения")
                return
            finally:
                self.root.config(cursor="")

            self.ds = ds
            self.results = []
            self.by_label = {}
            self.machine_info = None
            self.lengths = {}
            self.selected = set()
            self.iid_to_label = {}

            for b in ds.blocks:
                if b.label != ds.machine_label:
                    self.selected.add(b.label)

            file_lengths = [b.length for b in ds.blocks
                            if b.label != ds.machine_label and math.isfinite(b.length) and b.length > 0]
            if file_lengths:
                base_len = float(np.median(file_lengths))
                self.e_len.delete(0, "end")
                self.e_len.insert(0, f"{base_len:g}")
            for b in ds.blocks:
                if b.label != ds.machine_label:
                    self.lengths[b.label] = (b.length if math.isfinite(b.length) and b.length > 0
                                             else self.default_length())

            self.populate()
            self.txt.configure(state="normal")
            self.txt.delete("1.0", "end")
            self.txt.insert("end", "Файл загружен. Нажмите «Рассчитать».\n\n" + build_file_info(ds))
            self.txt.configure(state="disabled")
            self.file_info.configure(
                text=f"Серия {ds.series}: {len(self.lengths)} образцов"
                     + (" + образец 0" if ds.machine_label else ""))
            self.status.configure(text="Файл загружен. Нажмите «Рассчитать».")
            self.btn_xlsx.configure(state="disabled")
            self.btn_png.configure(state="disabled")
            self.clear_plots()

        def default_length(self) -> float:
            v = parse_number(self.e_len.get())
            return v if math.isfinite(v) and v > 0 else DEFAULT_LENGTH

        def populate(self):
            self.tree.delete(*self.tree.get_children())
            ds = self.ds
            if ds is None:
                return
            for b in ds.blocks:
                if b.label == ds.machine_label:
                    continue
                iid = f"s{b.label}"
                self.iid_to_label[iid] = b.label
                mark = "☑" if b.label in self.selected else "☐"
                self.tree.insert("", "end", iid=iid, values=(
                    mark, b.label, fmt(b.area, 1), fmt(self.lengths.get(b.label), 2),
                    "—", "—", "—", "—", ""))
            if ds.machine_label:
                t0, p0, u0 = ds.data[ds.machine_label]
                iid = "s0"
                self.iid_to_label[iid] = ds.machine_label
                self.tree.insert("", "end", iid=iid, values=(
                    "—", ds.machine_label, "—", "—", fmt_int(float(p0.max()) if p0.size else NAN),
                    "—", "—", "—", "без образца (эталон)"), tags=("machine",))

        # ---------------- выбор ----------------

        def on_click(self, event):
            if self.tree.identify("region", event.x, event.y) != "cell":
                return
            col = self.tree.identify_column(event.x)
            iid = self.tree.identify_row(event.y)
            if not iid or col != "#1":
                return
            label = self.iid_to_label.get(iid)
            if label is None or label == (self.ds.machine_label if self.ds else None):
                return
            self.toggle(label)

        def toggle(self, label):
            if label in self.selected:
                self.selected.discard(label)
            else:
                self.selected.add(label)
            iid = f"s{label}"
            if self.tree.exists(iid):
                self.tree.set(iid, "sel", "☑" if label in self.selected else "☐")
            self.refresh()

        def toggle_many(self):
            sel = self.tree.selection()
            if not sel:
                return
            labels = [self.iid_to_label[i] for i in sel
                      if i in self.iid_to_label and i != "s0"
                      and self.iid_to_label[i] != (self.ds.machine_label if self.ds else None)]
            if not labels:
                return
            make_selected = any(l not in self.selected for l in labels)
            for l in labels:
                if make_selected:
                    self.selected.add(l)
                else:
                    self.selected.discard(l)
                iid = f"s{l}"
                if self.tree.exists(iid):
                    self.tree.set(iid, "sel", "☑" if l in self.selected else "☐")
            self.refresh()

        def select_all(self):
            if not self.ds:
                return
            self.selected = {b.label for b in self.ds.blocks if b.label != self.ds.machine_label}
            self._sync_marks()
            self.refresh()

        def select_none(self):
            self.selected.clear()
            self._sync_marks()
            self.refresh()

        def _sync_marks(self):
            for iid, label in self.iid_to_label.items():
                if iid == "s0":
                    continue
                if self.tree.exists(iid):
                    self.tree.set(iid, "sel", "☑" if label in self.selected else "☐")

        # ---------------- правка L ----------------

        def on_double(self, event):
            if self.tree.identify("region", event.x, event.y) != "cell":
                return
            col = self.tree.identify_column(event.x)
            if col != "#4":
                return
            iid = self.tree.identify_row(event.y)
            if not iid or iid == "s0":
                return
            bbox = self.tree.bbox(iid, col)
            if not bbox:
                return
            x, y, w, h = bbox
            if self._editing is not None:
                self._editing.destroy()
            entry = ttk.Entry(self.tree)
            entry.place(x=x, y=y, width=w, height=h)
            entry.insert(0, self.tree.set(iid, "length"))
            entry.select_range(0, "end")
            entry.focus_set()
            self._editing = entry

            def commit(_=None):
                txt = entry.get().replace(",", ".").strip()
                try:
                    val = float(txt)
                    if not (val > 0):
                        raise ValueError
                except ValueError:
                    entry.destroy()
                    self._editing = None
                    return
                label = self.iid_to_label.get(iid)
                if label is not None:
                    self.lengths[label] = val
                self.tree.set(iid, "length", f"{val:.2f}")
                entry.destroy()
                self._editing = None

            entry.bind("<Return>", commit)
            entry.bind("<FocusOut>", commit)
            entry.bind("<Escape>", lambda e: (entry.destroy(), setattr(self, "_editing", None)))

        # ---------------- расчёт ----------------

        def recalc(self):
            if not self.ds:
                messagebox.showwarning("Нет данных", "Сначала загрузите файл Excel.")
                return
            try:
                m = int(float(self.e_m.get().replace(",", ".")))
                ms = int(float(self.e_ms.get().replace(",", ".")))
                if m < 2:
                    raise ValueError("m должно быть ≥ 2")
                if ms < 0:
                    raise ValueError("сглаживание не может быть отрицательным")
                area_txt = self.e_area.get().replace(",", ".").strip()
                area_fb = float(area_txt) if area_txt else NAN
                if area_txt and not (area_fb > 0):
                    raise ValueError("площадь должна быть положительной")
            except ValueError as exc:
                messagebox.showerror("Параметры", f"Некорректное значение: {exc}")
                return

            params = {
                "length": self.default_length(),
                "m": m,
                "machine_smooth": ms,
                "use_machine": bool(self.var_machine.get()),
            }

            self.root.config(cursor="watch")
            self.status.configure(text="Расчёт…")
            self.root.update_idletasks()
            try:
                results, machine_info = compute_all(
                    self.ds, self.lengths, m, ms,
                    area_fallback=area_fb, use_machine=params["use_machine"])
            except Exception as exc:
                self.root.config(cursor="")
                messagebox.showerror("Ошибка расчёта", str(exc))
                self.status.configure(text="Ошибка расчёта")
                return
            finally:
                self.root.config(cursor="")

            self.results = results
            self.by_label = {r.label: r for r in results}
            self.machine_info = machine_info

            for r in results:
                iid = f"s{r.label}"
                if not self.tree.exists(iid):
                    continue
                if r.ok:
                    self.tree.item(iid, values=(
                        "☑" if r.label in self.selected else "☐", r.label,
                        fmt(r.area, 1), fmt(r.length, 2), fmt_int(r.pmax), fmt(r.sigma, 1),
                        fmt(r.e_mpa / 1000.0 if math.isfinite(r.e_mpa) else NAN, 2),
                        fmt_int(r.k_corr), r.note), tags=())
                else:
                    self.tree.item(iid, values=(
                        "☐", r.label, fmt(r.area, 1), fmt(r.length, 2),
                        "—", "—", "—", "—", r.error), tags=("error",))
                    self.selected.discard(r.label)
                    self.tree.set(iid, "sel", "☐")

            n_ok = sum(1 for r in results if r.ok)
            self.status.configure(text=f"Готово: {n_ok} из {len(results)} образцов рассчитано")
            self.btn_xlsx.configure(state="normal" if n_ok else "disabled")
            self.btn_png.configure(state="normal" if n_ok else "disabled")
            self.refresh()

        def selected_results(self):
            return [r for r in self.results if r.ok and r.label in self.selected]

        # ---------------- отображение ----------------

        def refresh(self):
            if not self.results:
                return
            sel = self.selected_results()
            params = {
                "length": self.default_length(),
                "m": int(parse_number(self.e_m.get()) if math.isfinite(parse_number(self.e_m.get())) else DEFAULT_M),
                "machine_smooth": int(parse_number(self.e_ms.get()) if math.isfinite(parse_number(self.e_ms.get())) else DEFAULT_MACHINE_SMOOTH),
                "use_machine": bool(self.var_machine.get()),
            }
            text = build_summary_text(self.ds, self.results, sel, params)
            self.txt.configure(state="normal")
            self.txt.delete("1.0", "end")
            self.txt.insert("end", text)
            self.txt.configure(state="disabled")
            self.draw_overview(sel)
            self.draw_detail()

        def clear_plots(self):
            for frame in (self.tab_over, self.tab_detail):
                for w in frame.winfo_children():
                    w.destroy()

        def _draw_into(self, frame, fig):
            for w in frame.winfo_children():
                w.destroy()
            canvas = FigureCanvasTkAgg(fig, master=frame)
            canvas.draw()
            widget = canvas.get_tk_widget()
            widget.configure(width=380, height=330)
            widget.pack(fill="both", expand=True)

        def draw_overview(self, sel):
            if not sel:
                for w in self.tab_over.winfo_children():
                    w.destroy()
                ttk.Label(self.tab_over, text="Образцы не выбраны").pack(pady=40)
                return
            self._draw_into(self.tab_over, figure_overview(sel, self.machine_info))

        def draw_detail(self):
            label = None
            sel_rows = self.tree.selection()
            if sel_rows:
                label = self.iid_to_label.get(sel_rows[0])
            if label is None or label == (self.ds.machine_label if self.ds else None):
                cand = self.selected_results()
                label = cand[0].label if cand else None
            r = self.by_label.get(label) if label else None
            if r is None or not r.ok:
                for w in self.tab_detail.winfo_children():
                    w.destroy()
                ttk.Label(self.tab_detail, text="Выберите рассчитанный образец в таблице").pack(pady=40)
                return
            self._draw_into(self.tab_detail, figure_sample(r))

        def on_select_row(self):
            if self.results:
                self.draw_detail()

        # ---------------- экспорт ----------------

        def _params(self):
            return {
                "length": self.default_length(),
                "m": int(parse_number(self.e_m.get()) if math.isfinite(parse_number(self.e_m.get())) else DEFAULT_M),
                "machine_smooth": int(parse_number(self.e_ms.get()) if math.isfinite(parse_number(self.e_ms.get())) else DEFAULT_MACHINE_SMOOTH),
                "use_machine": bool(self.var_machine.get()),
            }

        def export_xlsx(self):
            sel = self.selected_results()
            if not sel:
                messagebox.showwarning("Экспорт", "Не выбран ни один рассчитанный образец.")
                return
            default = os.path.join(os.path.dirname(self.ds.path),
                                   f"{self.ds.series}_результаты.xlsx")
            path = filedialog.asksaveasfilename(
                title="Сохранить результаты", defaultextension=".xlsx",
                initialdir=os.path.dirname(default), initialfile=os.path.basename(default),
                filetypes=[("Excel", "*.xlsx")])
            if not path:
                return
            try:
                export_excel(path, self.ds, self.results, sel, self._params(), self.machine_info)
            except Exception as exc:
                messagebox.showerror("Ошибка сохранения", f"{exc}\n\n{traceback.format_exc(limit=2)}")
                return
            messagebox.showinfo("Готово", f"Сохранено:\n{path}")

        def export_png(self):
            sel = self.selected_results()
            if not sel:
                messagebox.showwarning("Экспорт", "Не выбран ни один рассчитанный образец.")
                return
            folder = filedialog.askdirectory(title="Папка для графиков",
                                             initialdir=os.path.dirname(self.ds.path))
            if not folder:
                return
            try:
                base = re.sub(r"[^\w\-. ]+", "_", self.ds.series).strip() or "серия"
                fig = figure_overview(sel, self.machine_info)
                p1 = os.path.join(folder, f"{base}_общие_графики.png")
                fig.savefig(p1, dpi=150)
                paths = [p1]
                for r in sel:
                    p = os.path.join(folder, f"{base}_образец_{r.label}.png")
                    figure_sample(r).savefig(p, dpi=150)
                    paths.append(p)
            except Exception as exc:
                messagebox.showerror("Ошибка сохранения", str(exc))
                return
            messagebox.showinfo("Готово", "Сохранено файлов: %d\n%s" % (len(paths), folder))

    root = tk.Tk()
    try:
        style = ttk.Style()
        if "vista" in style.theme_names():
            style.theme_use("vista")
    except Exception:
        pass
    App(root)
    root.mainloop()


if __name__ == "__main__":
    if len(sys.argv) == 1:
        main_gui()
    else:
        main_cli()
